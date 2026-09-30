import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    # Premium Developer-Product Dark Color Palette
    BG_COLOR = RGBColor(11, 15, 25)         # #0B0F19 Deep slate/black
    PANEL_BG = RGBColor(20, 27, 45)         # #141B2D Dark card background
    CARD_BORDER = RGBColor(40, 53, 80)      # #283550 Subtle border
    ACCENT_EMERALD = RGBColor(16, 185, 129) # #10B981 Emerald
    ACCENT_CYAN = RGBColor(6, 182, 212)     # #06B6D4 Cyan
    TEXT_WHITE = RGBColor(248, 250, 252)    # #F8FAFC Primary text (High contrast)
    TEXT_LIGHT = RGBColor(226, 232, 240)    # #E2E8F0 Light slate (Very readable)
    TEXT_MUTED = RGBColor(180, 195, 215)    # #B4C3D7 Secondary body (Brightened)
    TEXT_DIM = RGBColor(148, 163, 184)      # #94A3B8 Dim/Footers
    COLOR_CRITICAL = RGBColor(248, 113, 113)# #F87171 Red
    COLOR_HIGH = RGBColor(251, 146, 60)     # #FB923C Orange
    COLOR_MEDIUM = RGBColor(250, 204, 21)   # #FACC15 Amber

    FONT_FAMILY = "Segoe UI"
    MONO_FONT = "Consolas"

    def apply_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_para(tf, text="", font_name=FONT_FAMILY, font_size=11, font_color=TEXT_WHITE, bold=False, italic=False, align=None):
        if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == "":
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        if align:
            p.alignment = align
        for r in p.runs:
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = font_color
        return p

    def add_header(slide, category, title, subtitle=None):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(8), Inches(0.35))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        add_para(tf, category.upper(), font_size=10, font_color=ACCENT_EMERALD, bold=True)

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        add_para(tf, title, font_size=24, font_color=TEXT_WHITE, bold=True)

        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.32), Inches(11), Inches(0.4))
            tf = sub_box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            add_para(tf, subtitle, font_size=12, font_color=TEXT_MUTED)

    def add_footer(slide, current_idx, total_slides=15):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
        tf = footer_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        add_para(tf, f"DeploySafe — AI-Powered Deployment Readiness Platform  |  Agentic AI (CSE3101) Midterm  |  Slide {current_idx} of {total_slides}", font_size=9, font_color=TEXT_DIM)

    def add_panel(slide, left, top, width, height, bg_color=PANEL_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    def set_notes(slide, script, question, answer):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = f"=== 30-60 SEC TALKING SCRIPT ===\n{script}\n\n=== LIKELY PROFESSOR QUESTION ===\nQ: {question}\n\n=== CONCISE ANSWER ===\nA: {answer}"

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s1)

    # Accent decorative badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(3.4), Inches(0.38))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(6, 78, 59) # Emerald 900
    badge.line.color.rgb = ACCENT_EMERALD
    add_para(badge.text_frame, "AGENTIC AI — CSE3101 MIDTERM", font_size=10, font_color=ACCENT_EMERALD, bold=True, align=PP_ALIGN.CENTER)

    # Project Title
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(10), Inches(1.1))
    add_para(t_box.text_frame, "DeploySafe", font_size=46, font_color=TEXT_WHITE, bold=True)

    sub_box = s1.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(10), Inches(0.6))
    add_para(sub_box.text_frame, "Autonomous Multi-Agent Deployment-Readiness & Risk Intelligence Platform", font_size=18, font_color=ACCENT_CYAN, bold=True)

    one_liner = s1.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(10), Inches(0.6))
    add_para(one_liner.text_frame, "“Identify potential deployment risks before release through specialized agents, grounded evidence, and deterministic risk scoring.”", font_size=13, font_color=TEXT_LIGHT, italic=True)

    # Metadata Panel
    add_panel(s1, Inches(0.8), Inches(4.7), Inches(11.733), Inches(1.8))
    meta_box = s1.shapes.add_textbox(Inches(1.1), Inches(4.85), Inches(11.1), Inches(1.5))
    mtf = meta_box.text_frame
    mtf.word_wrap = True
    add_para(mtf, "Course: Agentic AI — CSE3101  |  Project: DeploySafe  |  Status: Phases 1 to 4 Complete", font_size=13, font_color=TEXT_WHITE, bold=True)
    add_para(mtf, "Architecture: CrewAI Collaborative Flow  •  RAG: Chroma Vector DB + Local Embeddings  •  Scoring: Pure Mathematical Deterministic Engine", font_size=11, font_color=ACCENT_EMERALD)
    add_para(mtf, "Presenter / Team: Lead Product Engineer & System Architect  |  Review Date: October 2026", font_size=11, font_color=TEXT_MUTED)

    add_footer(s1, 1)
    set_notes(s1,
        "Good morning professor and fellow classmates. Today we are presenting DeploySafe, an autonomous multi-agent platform designed to evaluate software deployment readiness before code reaches production. Modern pipelines focus on automated execution, but DeploySafe introduces a dedicated reasoning layer that cross-correlates code, dependencies, containers, tests, and documentation.",
        "Why is this project suited for an Agentic AI course rather than a standard software engineering project?",
        "Because determining deployment readiness requires distributed reasoning across heterogeneous domains—code, containers, security, and documentation. No single static tool or prompt can synthesize these interactions. DeploySafe coordinates 8 autonomous CrewAI agents that invoke specialized tools, correlate facts with vector RAG, and undergo adversarial Critic validation."
    )

    # =========================================================================
    # SLIDE 2: The Deployment Problem
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s2)
    add_header(s2, "Context & Motivation", "The Deployment Problem", "Passing automated CI tests does not mean a repository is safe for production deployment.")

    # Traditional Flow Diagram
    add_panel(s2, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.5))
    flow_box = s2.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(11.333), Inches(1.2))
    ftf = flow_box.text_frame
    add_para(ftf, "TRADITIONAL DEPLOYMENT PIPELINE:", font_name=MONO_FONT, font_size=11, font_color=ACCENT_CYAN, bold=True)
    add_para(ftf, "Code Change  ──▶  Unit Tests  ──▶  CI Build  ──▶  Container Push  ──▶  Production Outage", font_name=MONO_FONT, font_size=13, font_color=TEXT_WHITE, bold=True)
    add_para(ftf, "The Pipeline Passed! Yet the service crashed immediately due to an unpinned dependency and a missing REDIS_URL in docker-compose.", font_size=11, font_color=COLOR_CRITICAL)

    # 3 Problem Dimension Cards
    dims = [
        ("Disjointed Silos", "Unit tests test function logic, linters test syntax, and Docker builds test images. No tool checks if documented environment variables match compose manifests.", COLOR_HIGH),
        ("Silent Drift Hazards", "Floating versions (e.g. 'requests' or 'python:latest') pass CI today, but break tomorrow when upstream repositories publish a breaking release.", COLOR_MEDIUM),
        ("Context Blindness", "Static tools lack semantic awareness of documentation. They cannot read README.md architecture guidelines to discover missing operational requirements.", ACCENT_CYAN)
    ]
    card_w = Inches(3.75)
    for i, (d_title, d_desc, d_col) in enumerate(dims):
        cx = Inches(0.8) + i * Inches(4.0)
        add_panel(s2, cx, Inches(3.55), card_w, Inches(3.1))
        c_box = s2.shapes.add_textbox(cx + Inches(0.25), Inches(3.75), card_w - Inches(0.5), Inches(2.7))
        ctf = c_box.text_frame
        ctf.word_wrap = True
        add_para(ctf, d_title, font_size=14, font_color=d_col, bold=True)
        add_para(ctf, d_desc, font_size=11, font_color=TEXT_LIGHT)

    add_footer(s2, 2)
    set_notes(s2,
        "Every engineer has experienced a deployment failure where all unit tests passed green in GitHub Actions, but the service immediately crashed in staging or production. Traditional CI/CD verifies that individual components compile and unit tests pass, but it does not evaluate deployment context—such as unpinned base images, hardcoded endpoints, or documented environment prerequisites.",
        "Isn't this already solved by comprehensive integration tests and staging environments?",
        "Integration tests only validate paths developers anticipated and scripted. They do not catch operational configuration drift, unpinned supply-chain dependencies, documentation/manifest mismatches, or hardcoded local endpoints before containerization."
    )

    # =========================================================================
    # SLIDE 3: Why CI/CD is Not Enough
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s3)
    add_header(s3, "Product Positioning", "CI/CD Executes. DeploySafe Analyzes.", "DeploySafe does not replace CI/CD pipelines; it augments them with an AI reasoning and risk evaluation gate.")

    col_w = Inches(5.7)

    # Column 1: CI/CD
    add_panel(s3, Inches(0.8), Inches(1.8), col_w, Inches(4.9))
    ci_box = s3.shapes.add_textbox(Inches(1.1), Inches(2.0), col_w - Inches(0.6), Inches(4.5))
    tf1 = ci_box.text_frame
    tf1.word_wrap = True
    add_para(tf1, "TRADITIONAL CI/CD (EXECUTION ENGINE)", font_name=MONO_FONT, font_size=12, font_color=COLOR_HIGH, bold=True)

    items_ci = [
        ("Binary Success/Failure:", "Only checks if exit codes equal 0."),
        ("Predefined Scripted Checks:", "Only runs what engineers explicitly configured in YAML."),
        ("Isolated Job Boundaries:", "Test runner has zero awareness of Dockerfile or manifest changes."),
        ("Silent on Systemic Drift:", "Ignores unpinned libraries as long as pip install finishes."),
        ("No Evidence Explanation:", "Emits thousands of lines of raw terminal logs upon failure.")
    ]
    for h, b in items_ci:
        add_para(tf1, f"• {h} {b}", font_size=11, font_color=TEXT_LIGHT)

    # Column 2: DeploySafe
    add_panel(s3, Inches(6.833), Inches(1.8), col_w, Inches(4.9))
    ds_box = s3.shapes.add_textbox(Inches(7.133), Inches(2.0), col_w - Inches(0.6), Inches(4.5))
    tf2 = ds_box.text_frame
    tf2.word_wrap = True
    add_para(tf2, "DEPLOYSAFE (INTELLIGENCE & RISK LAYER)", font_name=MONO_FONT, font_size=12, font_color=ACCENT_EMERALD, bold=True)

    items_ds = [
        ("Semantic Risk Synthesis:", "Understands deployment intent across code, config, and docs."),
        ("Cross-Domain Correlation:", "Checks if environment variables in README exist in docker-compose."),
        ("Observable Evidence Contracts:", "Every finding links to exact file paths and 1-indexed lines."),
        ("Anti-Hallucination Critic:", "Adversarially validates physical file existence on disk."),
        ("Deterministic Risk Scoring:", "Computes bounded 0–100 scores without LLM mathematical drift.")
    ]
    for h, b in items_ds:
        add_para(tf2, f"✔ {h} {b}", font_size=11, font_color=TEXT_WHITE)

    add_footer(s3, 3)
    set_notes(s3,
        "It is vital to understand that DeploySafe is not competing with GitHub Actions or Jenkins. CI/CD is an execution engine that executes commands. DeploySafe is an analytical intelligence layer that reasons over the code, configuration, dependencies, and documentation. CI/CD asks 'did this command exit 0?'; DeploySafe asks 'is this change safe to release?'",
        "Could this be implemented as a SonarQube plugin or linter rule?",
        "No. Linters analyze single files using static AST patterns. They cannot correlate natural language documentation in README.md against docker-compose environment blocks, nor can they deduce missing regression coverage across interdependent service modules."
    )

    # =========================================================================
    # SLIDE 4: Our Solution - Meet DeploySafe
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s4)
    add_header(s4, "Solution Overview", "Meet DeploySafe", "A multi-agent intelligence layer that inspects repositories and delivers explainable deployment assessments.")

    # 3 Column layout
    add_panel(s4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.9))
    in_box = s4.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(3.2), Inches(4.5))
    tf1 = in_box.text_frame
    tf1.word_wrap = True
    add_para(tf1, "1. REPOSITORY SIGNALS", font_name=MONO_FONT, font_size=12, font_color=ACCENT_CYAN, bold=True)

    signals = [
        ("Source Code", "Git diffs, AST changes, endpoint calls, exception handlers."),
        ("Dependencies", "requirements.txt, package.json, lockfiles, unpinned packages."),
        ("Configuration", "Dockerfiles, compose services, environment configs."),
        ("Test Suite", "Test file presence, pytest runs, inferred coverage gaps."),
        ("Documentation", "README, architecture runbooks, required variables.")
    ]
    for s_name, s_det in signals:
        add_para(tf1, f"• {s_name}: {s_det}", font_size=10.5, font_color=TEXT_LIGHT)

    # Middle: Multi-Agent Engine
    add_panel(s4, Inches(4.7), Inches(1.8), Inches(3.9), Inches(4.9))
    mid_box = s4.shapes.add_textbox(Inches(4.9), Inches(2.0), Inches(3.5), Inches(4.5))
    tf2 = mid_box.text_frame
    tf2.word_wrap = True
    add_para(tf2, "2. DEPLOYSAFE ENGINE", font_name=MONO_FONT, font_size=12, font_color=ACCENT_EMERALD, bold=True)

    engine_items = [
        ("Orchestrator Flow", "Analyzes tech stack & dynamically plans tasks."),
        ("6 Specialized Agents", "Code, Dependencies, Config, Tests, RAG, Security."),
        ("Chroma Vector RAG", "Retrieves documentation and deployment rules."),
        ("Adversarial Critic", "Validates evidence and discards hallucinations."),
        ("Deduplication Service", "Merges multi-agent duplicate findings."),
        ("Deterministic Risk Engine", "Computes bounded mathematical 0–100 score.")
    ]
    for e_name, e_det in engine_items:
        add_para(tf2, f"▶ {e_name}: {e_det}", font_size=10.5, font_color=TEXT_WHITE)

    # Right: Actionable Output
    add_panel(s4, Inches(8.9), Inches(1.8), Inches(3.633), Inches(4.9))
    out_box = s4.shapes.add_textbox(Inches(9.1), Inches(2.0), Inches(3.2), Inches(4.5))
    tf3 = out_box.text_frame
    tf3.word_wrap = True
    add_para(tf3, "3. AUDITABLE OUTPUTS", font_name=MONO_FONT, font_size=12, font_color=COLOR_HIGH, bold=True)

    outputs = [
        ("Deterministic Score", "0–100 score + explainable risk level (e.g. CRITICAL)."),
        ("Validated Findings", "Ranked hazards with severity and confidence."),
        ("Concrete Evidence", "Direct file paths, line ranges, and verified code."),
        ("Explainable Rationale", "Answers: What? Why? Where? and Blast Radius."),
        ("Actionable Fixes", "Specific remediation advice, not generic platitudes."),
        ("Multi-Format Reports", "Interactive HTML, Markdown docs, and JSON contracts.")
    ]
    for o_name, o_det in outputs:
        add_para(tf3, f"✔ {o_name}: {o_det}", font_size=10.5, font_color=TEXT_WHITE)

    add_footer(s4, 4)
    set_notes(s4,
        "Here is the high-level system concept. DeploySafe ingests the complete set of repository signals—source code, manifests, Docker files, test suites, and documentation. These signals feed into the DeploySafe engine where specialized agents analyze their respective domains. Finally, the system produces actionable, evidence-backed outputs: a deterministic score, verified findings, and multi-format reports.",
        "Why is documentation included as a core repository signal?",
        "Because documentation represents human architectural intent. When a developer writes in README.md that REDIS_URL is required in production, that is a contractual deployment requirement. If docker-compose.yml fails to configure REDIS_URL, that is a high-risk configuration mismatch that traditional linters completely miss."
    )

    # =========================================================================
    # SLIDE 5: How DeploySafe Works (Pipeline Flow)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s5)
    add_header(s5, "Processing Pipeline", "From Repository to Risk Insight", "The end-to-end execution flow implemented across Phases 1 to 4.")

    steps = [
        ("1. Ingestion", "Clones git repo or local directory. Redacts secrets immediately.", ACCENT_CYAN),
        ("2. Discovery", "Extracts AST, manifests, Dockerfiles, and test suites.", ACCENT_CYAN),
        ("3. Vector RAG", "Chunks documentation into Chroma DB with semantic embeddings.", ACCENT_EMERALD),
        ("4. Agents", "Parallel execution across 6 specialized domain agents.", ACCENT_EMERALD),
        ("5. Critic Audit", "Adversarially checks physical disk files & line ranges.", COLOR_HIGH),
        ("6. Deduplication", "Clusters multi-agent overlap & elects canonical findings.", COLOR_HIGH),
        ("7. Risk Engine", "Computes bounded 0–100 score via pure deterministic math.", COLOR_CRITICAL)
    ]
    step_w = Inches(1.55)
    for i, (s_title, s_desc, s_color) in enumerate(steps):
        sx = Inches(0.8) + i * Inches(1.68)
        add_panel(s5, sx, Inches(2.0), step_w, Inches(3.2))
        s_box = s5.shapes.add_textbox(sx + Inches(0.12), Inches(2.15), step_w - Inches(0.24), Inches(2.9))
        stf = s_box.text_frame
        stf.word_wrap = True
        add_para(stf, s_title, font_size=12, font_color=s_color, bold=True)
        add_para(stf, s_desc, font_size=10, font_color=TEXT_LIGHT)

    # Bottom Pipeline Guarantees
    add_panel(s5, Inches(0.8), Inches(5.4), Inches(11.733), Inches(1.3))
    guar_box = s5.shapes.add_textbox(Inches(1.0), Inches(5.5), Inches(11.333), Inches(1.1))
    gtf = guar_box.text_frame
    add_para(gtf, "CORE PIPELINE GUARANTEES IMPLEMENTED IN CODE:", font_name=MONO_FONT, font_size=11, font_color=ACCENT_EMERALD, bold=True)
    add_para(gtf, "• Zero-Hallucination Gate: If a cited file does not physically exist on disk, the Critic Agent rejects it immediately.", font_size=10.5, font_color=TEXT_WHITE)
    add_para(gtf, "• Pure Math Risk Score: The LLM produces observations; the deterministic Risk Engine computes points. Same findings = exact same score.", font_size=10.5, font_color=TEXT_WHITE)

    add_footer(s5, 5)
    set_notes(s5,
        "This slide illustrates the 7-step pipeline implemented in DeploySafe. Notice the distinct boundary between Steps 4 and 7. The LLM agents operate in Step 4 to identify candidate hazards. But before those hazards affect the score, Step 5 validates physical file existence, Step 6 deduplicates multi-agent findings, and Step 7 applies a mathematical scoring formula.",
        "What happens if an LLM agent generates a finding for a file that was deleted or doesn't exist?",
        "Step 5, the Critic Agent, catches it. The Critic executes a physical filesystem stat check using the exact repository snapshot. If the file is not on disk, the finding is assigned status REJECTED, tagged as a hallucination, and contributes 0 points to the risk score."
    )

    # =========================================================================
    # SLIDE 6: System Architecture
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s6)
    add_header(s6, "System Design", "DeploySafe Technical Architecture", "Decoupling deterministic tools from agentic reasoning and evidence-backed synthesis.")

    # Top: React UI
    add_panel(s6, Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.65), bg_color=PANEL_BG, border_color=ACCENT_CYAN)
    ui_lbl = s6.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.333), Inches(0.45))
    add_para(ui_lbl.text_frame, "WEB UI: React 18  •  TypeScript  •  TailwindCSS  •  Live Telemetry  •  Evidence Inspector  •  Multi-Format Exports", font_name=MONO_FONT, font_size=11, font_color=ACCENT_CYAN, align=PP_ALIGN.CENTER)

    # Middle Gateway
    add_panel(s6, Inches(2.0), Inches(2.65), Inches(9.333), Inches(0.6), bg_color=PANEL_BG, border_color=ACCENT_EMERALD)
    gw_lbl = s6.shapes.add_textbox(Inches(2.2), Inches(2.72), Inches(8.933), Inches(0.45))
    add_para(gw_lbl.text_frame, "FASTAPI BACKEND GATEWAY: REST APIs (/dashboard, /runs, /findings, /analysis, /flow)", font_name=MONO_FONT, font_size=11, font_color=ACCENT_EMERALD, align=PP_ALIGN.CENTER)

    # Middle 3 columns: Agents, RAG, Tools
    sub_y = Inches(3.45)
    sub_h = Inches(2.0)
    col_w3 = Inches(3.75)

    # Box 1: Agents
    add_panel(s6, Inches(0.8), sub_y, col_w3, sub_h)
    ab_box = s6.shapes.add_textbox(Inches(0.95), sub_y + Inches(0.12), col_w3 - Inches(0.3), sub_h - Inches(0.24))
    atf = ab_box.text_frame
    atf.word_wrap = True
    add_para(atf, "CREWAI AGENTS (Parallel)", font_name=MONO_FONT, font_size=11, font_color=ACCENT_EMERALD, bold=True)
    for a in ["Code Analysis Agent", "Dependency Agent", "Configuration Agent", "Test Agent", "Security Agent"]:
        add_para(atf, f"• {a}", font_size=10, font_color=TEXT_WHITE)

    # Box 2: RAG
    add_panel(s6, Inches(4.791), sub_y, col_w3, sub_h)
    rb_box = s6.shapes.add_textbox(Inches(4.941), sub_y + Inches(0.12), col_w3 - Inches(0.3), sub_h - Inches(0.24))
    rtf = rb_box.text_frame
    rtf.word_wrap = True
    add_para(rtf, "KNOWLEDGE / RAG", font_name=MONO_FONT, font_size=11, font_color=ACCENT_CYAN, bold=True)
    for r in ["Documentation RAG Agent", "Chroma Vector Database", "Local Deterministic Embeddings", "Heading/Function Chunking", "Semantic Cosine Search"]:
        add_para(rtf, f"• {r}", font_size=10, font_color=TEXT_WHITE)

    # Box 3: Tools & Sandboxing
    add_panel(s6, Inches(8.783), sub_y, col_w3, sub_h)
    tb_box = s6.shapes.add_textbox(Inches(8.933), sub_y + Inches(0.12), col_w3 - Inches(0.3), sub_h - Inches(0.24))
    ttf = tb_box.text_frame
    ttf.word_wrap = True
    add_para(ttf, "TYPED TOOLS & SANDBOX", font_name=MONO_FONT, font_size=11, font_color=COLOR_HIGH, bold=True)
    for t in ["Sandboxed Safe Executor", "Binary Allowlists (pytest, git)", "Path Traversal Defense", "Entropy Secret Redaction", "AST & Manifest Parsers"]:
        add_para(ttf, f"• {t}", font_size=10, font_color=TEXT_WHITE)

    # Bottom: Critic + Dedup + Risk Engine
    add_panel(s6, Inches(0.8), Inches(5.65), Inches(11.733), Inches(1.15), bg_color=PANEL_BG, border_color=COLOR_CRITICAL)
    crit_box = s6.shapes.add_textbox(Inches(1.0), Inches(5.72), Inches(11.333), Inches(1.0))
    ctf = crit_box.text_frame
    add_para(ctf, "VERIFICATION & DETERMINISTIC SCORING ENGINE:", font_name=MONO_FONT, font_size=11, font_color=COLOR_CRITICAL, bold=True)
    add_para(ctf, "Candidate Findings ──▶ Critic Agent (Grounding) ──▶ Deduplication Service ──▶ RiskEngine (0–100 Bounded Score)", font_name=MONO_FONT, font_size=10.5, font_color=TEXT_WHITE, bold=True)

    add_footer(s6, 6)
    set_notes(s6,
        "Here is the complete Phase 4 architectural topology. The React frontend interacts with a versioned FastAPI gateway. The gateway invokes the CrewAI Flow Orchestrator, which dispatches specialized agents, vector RAG retrieval, and sandboxed deterministic tools. Crucially, all agent discoveries converge on the Critic Agent for grounding checks, pass through deduplication, and are scored by the mathematical Risk Engine.",
        "How is CrewAI integrated here?",
        "We implement DeploySafeAnalysisFlow using CrewAI's state flow abstraction. The flow handles repository context preparation, dynamic planning based on detected files, parallel agent dispatch via thread pools, and structured Pydantic state persistence."
    )

    # =========================================================================
    # SLIDE 7: Why Multi-Agent?
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s7)
    add_header(s7, "Agentic Architecture", "Why Multi-Agent? Specialized Domain Delegation", "A single monolithic prompt suffers from context dilution, hallucinations, and tool overload.")

    # 8 Agents Grid (2 rows x 4 cols)
    agents_data = [
        ("Code Analysis Agent", "Inspects source files for hardcoded localhost endpoints, swallowing exception handlers, and unsafe environment parsing.", ACCENT_CYAN),
        ("Dependency Agent", "Parses requirements.txt and package.json to identify unpinned dependencies and missing lockfiles.", ACCENT_EMERALD),
        ("Configuration Agent", "Audits Dockerfiles and compose manifests for unpinned base tags (python:latest) and config drift.", ACCENT_EMERALD),
        ("Test Analysis Agent", "Executes sandboxed pytest and strictly separates OBSERVED failures from INFERRED missing coverage.", COLOR_HIGH),
        ("Documentation RAG", "Queries Chroma vector store to extract documented mandatory requirements (e.g. REDIS_URL) & verifies compose blocks.", ACCENT_CYAN),
        ("Security Agent", "Runs static vulnerability scans, detects credentials, and masks secrets (***REDACTED_SECRET***).", COLOR_CRITICAL),
        ("Critic Gatekeeper", "Verifies that cited files physically exist on disk, validates line numbers, and rejects hallucinations.", COLOR_CRITICAL),
        ("Report Agent", "Synthesizes audited findings, compiles category distributions, and formats executive intelligence reports.", COLOR_HIGH)
    ]
    card_w4 = Inches(2.78)
    card_h4 = Inches(2.3)

    for idx, (a_name, a_desc, a_col) in enumerate(agents_data):
        row = idx // 4
        col = idx % 4
        ax = Inches(0.8) + col * Inches(2.98)
        ay = Inches(1.8) + row * Inches(2.5)

        add_panel(s7, ax, ay, card_w4, card_h4)
        abox = s7.shapes.add_textbox(ax + Inches(0.18), ay + Inches(0.18), card_w4 - Inches(0.36), card_h4 - Inches(0.36))
        atf = abox.text_frame
        atf.word_wrap = True
        add_para(atf, a_name, font_size=12, font_color=a_col, bold=True)
        add_para(atf, a_desc, font_size=10, font_color=TEXT_LIGHT)

    add_footer(s7, 7)
    set_notes(s7,
        "Why use 8 specialized agents instead of asking one large language model to do everything? In production repositories, monolithic prompts suffer from severe attention decay. By dedicating specialized agents with isolated system prompts and restricted tool access, each agent focuses deeply on its domain: Code, Dependencies, Configuration, Tests, RAG, and Security. Finally, the Critic agent acts as an independent adversarial auditor.",
        "Do these agents run sequentially or in parallel?",
        "The Orchestrator flow executes the independent domain agents concurrently using a ThreadPoolExecutor in Python. The Critic Agent and Report Agent run sequentially after all domain candidate findings have been collected."
    )

    # =========================================================================
    # SLIDE 8: RAG — How DeploySafe Understands the Repository
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s8)
    add_header(s8, "Knowledge Retrieval", "How DeploySafe Understands the Repository (RAG)", "Grounding agent reasoning in actual repository documentation and architecture runbooks.")

    # Left: RAG Pipeline
    add_panel(s8, Inches(0.8), Inches(1.8), Inches(4.5), Inches(4.9))
    rpipe_box = s8.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.5))
    rtf = rpipe_box.text_frame
    rtf.word_wrap = True
    add_para(rtf, "VECTOR RAG PIPELINE (PHASE 2)", font_name=MONO_FONT, font_size=12, font_color=ACCENT_CYAN, bold=True)

    r_steps = [
        ("File Discovery:", "Discovers markdown, documentation, and config files with path safety checks."),
        ("Semantic Chunking:", "Splits markdown along headings (H1/H2) and code along function boundaries."),
        ("Deterministic Embeddings:", "Chroma vector database indexed with 384-d normalized embeddings."),
        ("Cosine Similarity Retrieval:", "Retriever returns top-k chunks with similarity scores and line provenance.")
    ]
    for h, b in r_steps:
        add_para(rtf, f"• {h} {b}", font_size=11, font_color=TEXT_LIGHT)

    # Right: Concrete Real Example
    add_panel(s8, Inches(5.6), Inches(1.8), Inches(6.933), Inches(4.9))
    ex_box = s8.shapes.add_textbox(Inches(5.8), Inches(2.0), Inches(6.5), Inches(4.5))
    etf = ex_box.text_frame
    etf.word_wrap = True
    add_para(etf, "REAL DEMO EXAMPLE: DOCUMENTATION / CONFIG MISMATCH", font_name=MONO_FONT, font_size=12, font_color=ACCENT_EMERALD, bold=True)
    add_para(etf, "1. Agent Search Query:\n\"What required environment variables and services are documented for deployment?\"", font_name=MONO_FONT, font_size=10, font_color=ACCENT_CYAN)
    add_para(etf, "2. Chroma Vector Retrieval (README.md Lines 5–9):\n\"## Architecture & Production Deployment Requirements\n- Mandatory Requirement: Production deployment requires REDIS_URL for distributed token storage.\"", font_name=MONO_FONT, font_size=10, font_color=TEXT_LIGHT)
    add_para(etf, "3. Agent Cross-Correlation with docker-compose.yml:\nFound: services.web.environment defines ENVIRONMENT and PORT, but REDIS_URL is completely missing.", font_name=MONO_FONT, font_size=10, font_color=COLOR_HIGH)
    add_para(etf, "4. Resulting Finding (F-RAG-001):\nSeverity: HIGH  •  Score Contribution: +25 pts  •  Validated by Critic.", font_name=MONO_FONT, font_size=10.5, font_color=ACCENT_EMERALD, bold=True)

    add_footer(s8, 8)
    set_notes(s8,
        "Slide 8 shows our actual RAG implementation built in Phase 2 and leveraged in Phase 3 by the Documentation RAG agent. On the right is a real query executed on our demo repository. The RAG agent searched for deployment prerequisites, retrieved lines 5-9 of README.md stating that REDIS_URL is mandatory, and cross-examined docker-compose.yml to discover that REDIS_URL was missing. That is cross-domain intelligence.",
        "Why use local embeddings instead of relying entirely on external API embeddings?",
        "We implemented an LLM and embedding factory abstraction. While DeploySafe supports Gemini embeddings, our local embedding provider operates fully offline and deterministically. This guarantees that test suites run in 6 seconds with zero API rate limits, zero costs, and 100% test reproducibility."
    )

    # =========================================================================
    # SLIDE 9: Tools + Grounded Evidence
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s9)
    add_header(s9, "Evidence Grounding", "Agents Don't Just Guess: Typed Tool Contracts", "Every finding must be backed by traceable, observable repository evidence.")

    # Left: Tool Layer
    add_panel(s9, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.9))
    t_box = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(3.2), Inches(4.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    add_para(tf1, "DETERMINISTIC TOOLS", font_name=MONO_FONT, font_size=12, font_color=ACCENT_CYAN, bold=True)

    tools_list = [
        ("Repository Tools:", "File slicing, git diff parsing, line-number mapping."),
        ("Manifest Tools:", "requirements.txt, package.json AST parsers."),
        ("Safe Executor:", "Subprocess execution restricted to binary allowlist (pytest, git)."),
        ("Security Scanner:", "Regex + entropy detector with automated secret masking."),
        ("RAG Retriever:", "Chroma semantic vector query tool.")
    ]
    for th, tb in tools_list:
        add_para(tf1, f"• {th} {tb}", font_size=10.5, font_color=TEXT_LIGHT)

    # Right: Concrete Finding Evidence Schema
    add_panel(s9, Inches(4.7), Inches(1.8), Inches(7.833), Inches(4.9))
    e_box = s9.shapes.add_textbox(Inches(4.9), Inches(2.0), Inches(7.4), Inches(4.5))
    tf2 = e_box.text_frame
    tf2.word_wrap = True
    add_para(tf2, "STRUCTURED EVIDENCE CONTRACT IN DEPLOYSAFE", font_name=MONO_FONT, font_size=12, font_color=ACCENT_EMERALD, bold=True)

    code_snippet = """Finding ID:    F-SEC-002 (Potential Hardcoded Secret)
File:          config/production.py : Line 5
Detected By:   SecurityAnalysisAgent (via deploy_safe_security_engine)
Status:        VALIDATED (Grounded on disk, secret masked)

Evidence:
  L4:  # Production database connection
  L5:  DATABASE_URI = "***REDACTED_SECRET***:5432/main"

Impact:        Exposed credentials will be baked into container images.
Recommendation: Extract into DATABASE_URL environment variable and rotate immediately."""

    add_para(tf2, code_snippet, font_name=MONO_FONT, font_size=9.5, font_color=TEXT_WHITE)
    add_para(tf2, "✔ Secret Masking Guarantee: Raw passwords and tokens are never emitted into LLM prompts, context windows, or stored reports.", font_size=10.5, font_color=ACCENT_EMERALD, bold=True)

    add_footer(s9, 9)
    set_notes(s9,
        "A critical vulnerability in LLM-based developer tools is the lack of verifiable evidence. In DeploySafe, agents cannot make vague statements like 'we think security is bad.' Every finding must fulfill a strict Pydantic Evidence schema: the exact repository, exact file path, exact 1-indexed start and end lines, and a verifiable code snippet. Furthermore, our security scanner automatically redacts secrets so raw tokens never leak.",
        "What happens if a tool produces invalid output or errors?",
        "All tools run inside defensive exception wrappers. If a tool fails or attempts path traversal outside the repository sandbox, the tool returns a structured error object. The agent logs the failure in its telemetry and continues without crashing the pipeline."
    )

    # =========================================================================
    # SLIDE 10: Deterministic Risk Engine
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s10)
    add_header(s10, "Risk Scoring", "From Findings to Deterministic Risk", "The LLM does not calculate the risk score. DeploySafe uses a strictly reproducible mathematical engine.")

    # Left: Formula & Weights
    add_panel(s10, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.9))
    f_box = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.6), Inches(4.5))
    tf1 = f_box.text_frame
    tf1.word_wrap = True
    add_para(tf1, "THE DETERMINISTIC RISK FORMULA", font_name=MONO_FONT, font_size=12, font_color=ACCENT_EMERALD, bold=True)
    add_para(tf1, "Score = min(100.0, round( ∑ (Weight × Confidence × Impact) ))", font_name=MONO_FONT, font_size=11, font_color=ACCENT_CYAN, bold=True)
    add_para(tf1, "Project-Defined Severity Weights:", font_size=11, font_color=TEXT_WHITE, bold=True)
    add_para(tf1, "• CRITICAL = 35.0 pts (Blockers: exposed secrets, remote execution)\n• HIGH     = 25.0 pts (Severe drifts: unpinned images, missing env)\n• MEDIUM   = 15.0 pts (Warnings: missing critical service tests)\n• LOW      =  7.0 pts (Cosmetic: minor documentation gaps)", font_name=MONO_FONT, font_size=10, font_color=TEXT_LIGHT)
    add_para(tf1, "Multipliers: High = 1.0  •  Medium = 0.8  •  Low = 0.6\nExclusions: Duplicates, rejected, & dismissed findings = 0.0 pts.", font_size=10.5, font_color=TEXT_LIGHT)

    # Right: Risk Levels & Actual Demo Math
    add_panel(s10, Inches(7.1), Inches(1.8), Inches(5.433), Inches(4.9))
    r_box = s10.shapes.add_textbox(Inches(7.3), Inches(2.0), Inches(5.0), Inches(4.5))
    tf2 = r_box.text_frame
    tf2.word_wrap = True
    add_para(tf2, "RISK TIERS & ACTUAL DEMO SCORE", font_name=MONO_FONT, font_size=12, font_color=COLOR_HIGH, bold=True)

    tiers = [
        ("0 – 24", "LOW RISK", "Ready for automated deployment.", ACCENT_EMERALD),
        ("25 – 49", "MODERATE RISK", "Review recommended before release.", COLOR_MEDIUM),
        ("50 – 74", "HIGH RISK", "Severe drift; manual sign-off required.", COLOR_HIGH),
        ("75 – 100", "CRITICAL RISK", "Deployment strictly blocked.", COLOR_CRITICAL)
    ]
    for r_range, r_name, r_act, r_col in tiers:
        add_para(tf2, f"{r_range}: {r_name} — {r_act}", font_size=10, font_color=r_col, bold=True)

    add_para(tf2, "Actual Measured Demo Run (deploysafe-analysis-demo):\n• 2 Critical findings: 2 × 35.0 = 70.0 pts\n• 4 Validated High findings: 4 × 25.0 = 100.0 pts\n• Raw Score: 170.0 pts ──▶ Capped at 100.0 / 100 (CRITICAL RISK)", font_name=MONO_FONT, font_size=9.5, font_color=TEXT_WHITE)
    add_para(tf2, "*Note: Risk thresholds are DeploySafe project-defined rules, not a guarantee of deployment safety.", font_size=9, font_color=TEXT_DIM, italic=True)

    add_footer(s10, 10)
    set_notes(s10,
        "Slide 10 details our Phase 4 deterministic Risk Engine. If you ask an LLM 'what is the risk on a scale of 1 to 100?', it might answer 75 on Monday and 60 on Tuesday. In DeploySafe, the LLM only identifies candidate findings. The RiskEngine runs pure arithmetic with project-defined severity weights: Critical=35, High=25, Medium=15, Low=7. The same findings will always yield the exact same score.",
        "Why is the score capped at 100?",
        "Risk represents an operational threshold. If a repository contains two critical hardcoded database passwords and four severe configuration mismatches, the raw total is 170 points. Capping at 100 maps the metric cleanly to a standardized 0–100 scale where anything 75 or higher is classified as CRITICAL and blocks deployment."
    )

    # =========================================================================
    # SLIDE 11: Product UI Experience (REAL SCREENSHOTS)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s11)
    add_header(s11, "User Experience", "DeploySafe Product UI Experience", "Production-grade developer intelligence interface implemented in React and Tailwind.")

    img_w = Inches(5.6)
    img_h = Inches(2.35)

    dash_img_path = "docs/assets/screenshot_dashboard.png"
    if os.path.exists(dash_img_path):
        s11.shapes.add_picture(dash_img_path, Inches(0.8), Inches(1.8), width=img_w, height=img_h)
    add_panel(s11, Inches(0.8), Inches(4.15), img_w, Inches(0.4), bg_color=PANEL_BG, border_color=CARD_BORDER)
    l1 = s11.shapes.add_textbox(Inches(0.9), Inches(4.17), img_w - Inches(0.2), Inches(0.35))
    add_para(l1.text_frame, "Overview Dashboard: Risk Score Hero (100/100), Severity Breakdown & Formula Modal", font_size=9.5, font_color=ACCENT_CYAN)

    find_img_path = "docs/assets/screenshot_findings.png"
    if os.path.exists(find_img_path):
        s11.shapes.add_picture(find_img_path, Inches(6.933), Inches(1.8), width=img_w, height=img_h)
    add_panel(s11, Inches(6.933), Inches(4.15), img_w, Inches(0.4), bg_color=PANEL_BG, border_color=CARD_BORDER)
    l2 = s11.shapes.add_textbox(Inches(7.033), Inches(4.17), img_w - Inches(0.2), Inches(0.35))
    add_para(l2.text_frame, "Findings Directory: Search, Multi-Filter (Severity/Status), & Risk Score Contributions", font_size=9.5, font_color=ACCENT_EMERALD)

    runs_img_path = "docs/assets/screenshot_runs.png"
    if os.path.exists(runs_img_path):
        s11.shapes.add_picture(runs_img_path, Inches(0.8), Inches(4.65), width=img_w, height=img_h)
    add_panel(s11, Inches(0.8), Inches(6.6), img_w, Inches(0.35), bg_color=PANEL_BG, border_color=CARD_BORDER)
    l3 = s11.shapes.add_textbox(Inches(0.9), Inches(6.62), img_w - Inches(0.2), Inches(0.3))
    add_para(l3.text_frame, "Historical Runs: Duration, Findings Breakdown, Risk Level, & Agent Telemetry", font_size=9.5, font_color=COLOR_HIGH)

    rep_img_path = "docs/assets/screenshot_reports.png"
    if os.path.exists(rep_img_path):
        s11.shapes.add_picture(rep_img_path, Inches(6.933), Inches(4.65), width=img_w, height=img_h)
    add_panel(s11, Inches(6.933), Inches(6.6), img_w, Inches(0.35), bg_color=PANEL_BG, border_color=CARD_BORDER)
    l4 = s11.shapes.add_textbox(Inches(7.033), Inches(6.62), img_w - Inches(0.2), Inches(0.3))
    add_para(l4.text_frame, "Deployment Reports: Interactive HTML Preview, Documentation Markdown, & JSON Export", font_size=9.5, font_color=TEXT_WHITE)

    add_footer(s11, 11)
    set_notes(s11,
        "Here are authentic, unedited screenshots of the working DeploySafe web interface captured directly from our running Vite application. On the top left is the Overview Dashboard displaying our 100/100 CRITICAL risk hero. Top right is the Findings Directory showing live score contributions like +35 pts. Bottom left shows run history and execution telemetry. Bottom right shows our multi-format report exports.",
        "Did you use a component library like Tailwind or Material UI?",
        "We used Vanilla CSS and TailwindCSS for fine-grained typography and layout control, deliberately avoiding bulky pre-made component libraries to maintain high information density, fast rendering, and an authentic developer-security product aesthetic."
    )

    # =========================================================================
    # SLIDE 12: Live Demonstration Flow
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s12)
    add_header(s12, "Demonstration Guide", "What We Can Demonstrate Live", "A 7-step interactive walkthrough using our seeded demo repository.")

    demo_steps = [
        ("Step 1: Target Selection", "Select demo repository (deploysafe-analysis-demo) in the UI console.", "01"),
        ("Step 2: Start Analysis", "Trigger DeploySafe multi-agent analysis flow with a single click.", "02"),
        ("Step 3: Live Telemetry", "Observe 8 agents executing with live durations, tool calls, and counts.", "03"),
        ("Step 4: Vector RAG", "Watch RAG correlate README.md requirements against docker-compose.", "04"),
        ("Step 5: Critic & Dedup", "Critic audits physical disk paths; Deduplication clusters overlapping lines.", "05"),
        ("Step 6: Deterministic Score", "Inspect 100/100 score hero and open 'How is this calculated?' modal.", "06"),
        ("Step 7: Triage & Export", "Inspect code evidence, triage finding to DISMISSED, and export HTML.", "07")
    ]
    card_w7 = Inches(1.55)
    for i, (d_title, d_desc, d_num) in enumerate(demo_steps):
        dx = Inches(0.8) + i * Inches(1.68)
        add_panel(s12, dx, Inches(1.8), card_w7, Inches(4.9))

        np = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, dx + Inches(0.4), Inches(2.0), Inches(0.75), Inches(0.35))
        np.fill.solid()
        np.fill.fore_color.rgb = RGBColor(6, 78, 59)
        np.line.color.rgb = ACCENT_EMERALD
        add_para(np.text_frame, d_num, font_name=MONO_FONT, font_size=11, font_color=ACCENT_EMERALD, bold=True, align=PP_ALIGN.CENTER)

        dbox = s12.shapes.add_textbox(dx + Inches(0.1), Inches(2.5), card_w7 - Inches(0.2), Inches(4.0))
        dtf = dbox.text_frame
        dtf.word_wrap = True
        add_para(dtf, d_title, font_size=11, font_color=TEXT_WHITE, bold=True)
        add_para(dtf, d_desc, font_size=10, font_color=TEXT_LIGHT)

    add_footer(s12, 12)
    set_notes(s12,
        "During our live demo, we will walk the professor through these exact seven steps. We start by targeting our demo repository, trigger the analysis, observe the live agent activity timeline, watch RAG cross-correlate README requirements against compose manifests, verify how the Critic and Deduplication engines filter candidate findings, inspect the score explanation modal, and finally demonstrate status triage.",
        "How long does the live analysis take?",
        "When running in local diagnostic mode or with LLM caching, the full 8-agent analysis, Critic audit, vector RAG retrieval, deduplication, and risk calculation completes in approximately 1.5 to 2.5 seconds."
    )

    # =========================================================================
    # SLIDE 13: What We Have Completed (Phases 1–4)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s13)
    add_header(s13, "Implementation Status", "Current Implementation — Phases 1 to 4 Complete", "100% of Phase 1 through Phase 4 requirements from prd.md are implemented and verified.")

    p_w = Inches(2.8)
    phases = [
        ("PHASE 1: Foundation", [
            "✔ FastAPI Async Gateway",
            "✔ React + TypeScript Web Shell",
            "✔ Provider Factory (Gemini + Groq)",
            "✔ Offline Diagnostic Mock Mode",
            "✔ Pydantic Schemas & SQLite DB"
        ], ACCENT_CYAN),
        ("PHASE 2: Repository + RAG", [
            "✔ Repository Loader (Local + Git)",
            "✔ Safe File Discovery & AST",
            "✔ Markdown & Code Chunking",
            "✔ Chroma Persistent Vector DB",
            "✔ Semantic Cosine Retrieval"
        ], ACCENT_EMERALD),
        ("PHASE 3: Multi-Agent Engine", [
            "✔ CrewAI Flow Orchestrator",
            "✔ 6 Specialized Domain Agents",
            "✔ Sandboxed Tool Layer (Pytest, Git)",
            "✔ Adversarial Critic Gatekeeper",
            "✔ Line-Accurate Evidence Mapping"
        ], COLOR_HIGH),
        ("PHASE 4: Risk Engine & UI", [
            "✔ Deterministic RiskEngine (0–100)",
            "✔ Multi-Attribute Deduplication",
            "✔ Auditable Status Lifecycle",
            "✔ Professional React Dashboard",
            "✔ Multi-Format Reports (HTML/MD/JSON)"
        ], COLOR_CRITICAL)
    ]

    for i, (p_title, p_items, p_color) in enumerate(phases):
        px = Inches(0.8) + i * Inches(2.98)
        add_panel(s13, px, Inches(1.8), p_w, Inches(4.0))
        pbox = s13.shapes.add_textbox(px + Inches(0.15), Inches(1.95), p_w - Inches(0.3), Inches(3.7))
        ptf = pbox.text_frame
        ptf.word_wrap = True
        add_para(ptf, p_title, font_size=12, font_color=p_color, bold=True)
        for item in p_items:
            add_para(ptf, item, font_size=10, font_color=TEXT_WHITE)

    # Test Suite Banner
    add_panel(s13, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.7), bg_color=PANEL_BG, border_color=ACCENT_EMERALD)
    ts_box = s13.shapes.add_textbox(Inches(1.0), Inches(6.1), Inches(11.333), Inches(0.5))
    add_para(ts_box.text_frame, "AUTOMATED TEST SUITE STATUS: 62 of 62 Pytest Tests Passing (0 Failures)  •  Frontend Vite Build: 0 Errors in 2.38s", font_name=MONO_FONT, font_size=11, font_color=ACCENT_EMERALD, bold=True)

    add_footer(s13, 13)
    set_notes(s13,
        "This slide summarizes our rigorous engineering progress across all four completed phases. In Phase 1, we established the core FastAPI, React, and LLM abstraction. In Phase 2, we built the repository ingestion and Chroma vector store. In Phase 3, we implemented the 8 specialized CrewAI agents and tools. In Phase 4, we built the deterministic Risk Engine, deduplication service, and product UI. All 62 backend pytest tests pass 100%.",
        "How do you test the multi-agent system reliably without incurring large API fees?",
        "We implemented a test architecture that combines unit mocks with our offline diagnostic provider and local embedding generator. All 62 tests run deterministically in under 7 seconds without network dependency, allowing rapid CI validation."
    )

    # =========================================================================
    # SLIDE 14: What Comes Next (Roadmap)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s14)
    add_header(s14, "Future Roadmap", "What Comes Next — Phases 5 & 6", "Clear separation between what is working today and upcoming development phases.")

    col_w2 = Inches(5.7)

    # Left: Phase 5
    add_panel(s14, Inches(0.8), Inches(1.8), col_w2, Inches(4.9))
    p5_box = s14.shapes.add_textbox(Inches(1.1), Inches(2.0), col_w2 - Inches(0.6), Inches(4.5))
    tf1 = p5_box.text_frame
    tf1.word_wrap = True
    add_para(tf1, "PHASE 5: GITHUB AUTOMATION & EVALUATION (PLANNED)", font_name=MONO_FONT, font_size=12, font_color=COLOR_HIGH, bold=True)

    p5_items = [
        ("GitHub Action CI Workflow:", "Automated trigger on pull_request events (.github/workflows/deploysafe.yml)."),
        ("Automated PR Bot Comments:", "Posts structured risk assessment and blocker warnings directly on pull requests."),
        ("Benchmark Evaluation Dataset:", "Quantitative evaluation against curated repositories with known deployment flaws."),
        ("Safety & Prompt Hardening:", "Adversarial test suite evaluating defenses against prompt injection attacks."),
        ("Strict CI Gate Blocking:", "Exit code non-zero when Risk Score >= 75 (CRITICAL).")
    ]
    for h, b in p5_items:
        add_para(tf1, f"• {h} {b}", font_size=11, font_color=TEXT_LIGHT)

    # Right: Phase 6
    add_panel(s14, Inches(6.833), Inches(1.8), col_w2, Inches(4.9))
    p6_box = s14.shapes.add_textbox(Inches(7.133), Inches(2.0), col_w2 - Inches(0.6), Inches(4.5))
    tf2 = p6_box.text_frame
    tf2.word_wrap = True
    add_para(tf2, "PHASE 6: OBSERVABILITY & PRODUCTION READY (PLANNED)", font_name=MONO_FONT, font_size=12, font_color=ACCENT_CYAN, bold=True)

    p6_items = [
        ("Telemetry & Tracing:", "OpenTelemetry spans tracking LLM token costs, latencies, and agent execution paths."),
        ("Persistent Team Multi-Tenancy:", "User authentication, team workspaces, and repository access management."),
        ("Continuous Drift Monitoring:", "Periodic scheduled audits flagging upstream dependency deprecations."),
        ("Dockerized Distribution:", "Production multi-container Docker compose deployment for enterprise teams.")
    ]
    for h, b in p6_items:
        add_para(tf2, f"• {h} {b}", font_size=11, font_color=TEXT_LIGHT)

    add_footer(s14, 14)
    set_notes(s14,
        "We are strictly transparent about project boundaries. What we showed today is Phases 1 through 4. For the final phase of the course, Phase 5 will introduce automated GitHub Actions integration, PR comment bots, and empirical evaluation against a benchmark dataset. Phase 6 will introduce enterprise observability, token telemetry, and multi-tenant security.",
        "When will the GitHub PR bot be functional?",
        "Phase 5 begins immediately after midterm evaluation. The core backend APIs and report generation services built in Phase 4 already output the exact JSON contract needed by GitHub Actions."
    )

    # =========================================================================
    # SLIDE 15: Conclusion & Q&A
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s15)

    c_box = s15.shapes.add_textbox(Inches(1.5), Inches(1.4), Inches(10.333), Inches(3.0))
    ctf = c_box.text_frame
    ctf.word_wrap = True
    add_para(ctf, "DeploySafe", font_size=46, font_color=TEXT_WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_para(ctf, "“An AI analysis layer between code changes and deployment.”", font_size=18, font_color=ACCENT_EMERALD, italic=True, align=PP_ALIGN.CENTER)
    add_para(ctf, "Repository  ──▶  Agents  ──▶  Evidence  ──▶  Deterministic Risk  ──▶  Actionable Insight", font_name=MONO_FONT, font_size=13, font_color=ACCENT_CYAN, bold=True, align=PP_ALIGN.CENTER)

    # Thank you panel
    add_panel(s15, Inches(3.2), Inches(4.3), Inches(6.933), Inches(1.8))
    t_box = s15.shapes.add_textbox(Inches(3.4), Inches(4.45), Inches(6.533), Inches(1.5))
    ttf = t_box.text_frame
    ttf.word_wrap = True
    add_para(ttf, "Thank You!", font_size=24, font_color=TEXT_WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_para(ttf, "We are ready for questions and the live product demonstration.", font_size=13, font_color=ACCENT_EMERALD, align=PP_ALIGN.CENTER)
    add_para(ttf, "GitHub / Project Source: DeploySafe Core  •  Course: Agentic AI — CSE3101", font_size=11, font_color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    add_footer(s15, 15)
    set_notes(s15,
        "In conclusion, DeploySafe bridges the critical gap between traditional CI compilation and deployment confidence. By combining specialized CrewAI agents, Chroma vector RAG, anti-hallucination Critic validation, and a deterministic mathematical Risk Engine, DeploySafe gives teams true deployment intelligence. Thank you, and we welcome your questions.",
        "What is the single biggest technical takeaway from building DeploySafe?",
        "The biggest takeaway is that Agentic AI in production requires deterministic scaffolding. LLMs excel at semantic correlation and finding discovery across diverse files, but they must be grounded by strict tool contracts, validated by an adversarial Critic, and scored by a deterministic mathematical engine rather than relying on LLM intuition."
    )

    output_path = "DeploySafe_Midterm_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
